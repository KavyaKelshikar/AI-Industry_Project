"""
DOCX document extractor using python-docx with paragraph and table provenance.
"""

import logging
from pathlib import Path
from typing import List, Union
import docx

from src.document_processing.cleaner import clean_text
from src.document_processing.exceptions import (
    CorruptedFileError,
    EmptyDocumentError,
    ExtractorError,
)
from src.document_processing.extractor import (
    BaseExtractor,
    ExtractedDocument,
    ExtractedUnit,
    register_extractor,
)

logger = logging.getLogger(__name__)


class DOCXExtractor(BaseExtractor):
    """Extractor for Microsoft Word (.docx) documents capturing paragraphs and tables."""

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def extract(self, file_path: Union[str, Path]) -> ExtractedDocument:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"DOCX file not found: {path}")

        try:
            doc = docx.Document(path)
        except Exception as exc:
            raise CorruptedFileError(f"Failed to open/parse DOCX file {path.name}: {exc}", original_error=exc)

        units: List[ExtractedUnit] = []
        raw_texts: List[str] = []
        section_idx = 0

        # Extract paragraphs
        for p in doc.paragraphs:
            cleaned_p = clean_text(p.text)
            if cleaned_p:
                unit = ExtractedUnit(
                    text=cleaned_p,
                    page=None,
                    section_index=section_idx,
                    source=path.name,
                    metadata={"element_type": "paragraph", "style": p.style.name if p.style else None},
                )
                units.append(unit)
                raw_texts.append(cleaned_p)
                section_idx += 1

        # Extract tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [clean_text(cell.text) for cell in row.cells]
                row_str = " | ".join(filter(None, row_cells))
                if row_str:
                    table_rows.append(row_str)

            if table_rows:
                table_text = "\n".join(table_rows)
                unit = ExtractedUnit(
                    text=table_text,
                    page=None,
                    section_index=section_idx,
                    source=path.name,
                    metadata={"element_type": "table", "table_index": t_idx},
                )
                units.append(unit)
                raw_texts.append(table_text)
                section_idx += 1

        full_text = "\n\n".join(raw_texts).strip()
        if not full_text:
            raise EmptyDocumentError(f"DOCX document {path.name} contains no readable text")

        return ExtractedDocument(
            file_path=str(path.resolve()),
            file_name=path.name,
            file_type="docx",
            units=units,
            raw_text=full_text,
            total_pages=None,
            metadata={"source": path.name, "total_sections": len(units)},
        )


register_extractor(".docx", DOCXExtractor)
