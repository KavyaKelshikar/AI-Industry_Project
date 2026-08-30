"""
Tests for TXT, PDF, and DOCX text extractors and factory.
"""

import os
import tempfile
import docx
import fitz  # PyMuPDF
import pytest

from src.document_processing.exceptions import (
    CorruptedFileError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from src.document_processing.extractor import get_extractor
from src.document_processing.docx_extractor import DOCXExtractor
from src.document_processing.pdf_extractor import PDFExtractor
from src.document_processing.txt_extractor import TXTExtractor


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmp:
        yield tmp


class TestTXTExtractor:
    def test_extract_txt_success(self, temp_dir):
        file_path = os.path.join(temp_dir, "sample.txt")
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("Line 1 of sample text.\n\nLine 2 of sample text.")

        extractor = TXTExtractor()
        doc = extractor.extract(file_path)

        assert doc.file_name == "sample.txt"
        assert doc.file_type == "txt"
        assert len(doc.units) == 1
        assert "Line 1 of sample text." in doc.raw_text
        assert "Line 2 of sample text." in doc.raw_text
        assert doc.units[0].page == 1

    def test_extract_txt_empty_raises_empty_document_error(self, temp_dir):
        file_path = os.path.join(temp_dir, "empty.txt")
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("   \n\n   \t  ")

        extractor = TXTExtractor()
        with pytest.raises(EmptyDocumentError):
            extractor.extract(file_path)

    def test_extract_nonexistent_file_raises_not_found(self):
        extractor = TXTExtractor()
        with pytest.raises(FileNotFoundError):
            extractor.extract("non_existent_path_xyz.txt")


class TestPDFExtractor:
    def test_extract_pdf_multi_page_success(self, temp_dir):
        file_path = os.path.join(temp_dir, "multipage.pdf")
        # Create a synthetic 2-page PDF with fitz
        pdf_doc = fitz.open()
        p1 = pdf_doc.new_page()
        p1.insert_text((50, 50), "Page 1 Content: Industrial Safety Standard.")
        p2 = pdf_doc.new_page()
        p2.insert_text((50, 50), "Page 2 Content: Emergency Procedures.")
        pdf_doc.save(file_path)
        pdf_doc.close()

        extractor = PDFExtractor()
        doc = extractor.extract(file_path)

        assert doc.file_type == "pdf"
        assert doc.total_pages == 2
        assert len(doc.units) == 2
        assert doc.units[0].page == 1
        assert "Page 1 Content" in doc.units[0].text
        assert doc.units[1].page == 2
        assert "Page 2 Content" in doc.units[1].text

    def test_extract_empty_pdf_raises_empty_document_error(self, temp_dir):
        file_path = os.path.join(temp_dir, "blank.pdf")
        pdf_doc = fitz.open()
        pdf_doc.new_page()  # Blank page with no text
        pdf_doc.save(file_path)
        pdf_doc.close()

        extractor = PDFExtractor()
        with pytest.raises(EmptyDocumentError):
            extractor.extract(file_path)

    def test_extract_corrupted_pdf_raises_corrupted_error(self, temp_dir):
        file_path = os.path.join(temp_dir, "corrupt.pdf")
        with open(file_path, "wb") as f:
            f.write(b"NOT_A_VALID_PDF_HEADER_DATA_12345")

        extractor = PDFExtractor()
        with pytest.raises(CorruptedFileError):
            extractor.extract(file_path)


class TestDOCXExtractor:
    def test_extract_docx_paragraphs_and_tables(self, temp_dir):
        file_path = os.path.join(temp_dir, "sample.docx")
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1: Standard Operating Procedure")
        doc.add_paragraph("Paragraph 2: Maintenance Checklist")
        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Header A"
        table.rows[0].cells[1].text = "Header B"
        table.rows[1].cells[0].text = "Value 1"
        table.rows[1].cells[1].text = "Value 2"
        doc.save(file_path)

        extractor = DOCXExtractor()
        extracted = extractor.extract(file_path)

        assert extracted.file_type == "docx"
        assert len(extracted.units) == 3  # 2 paragraphs + 1 table
        assert "Standard Operating Procedure" in extracted.units[0].text
        assert "Maintenance Checklist" in extracted.units[1].text
        assert "Header A | Header B" in extracted.units[2].text

    def test_extract_empty_docx_raises_empty_document_error(self, temp_dir):
        file_path = os.path.join(temp_dir, "empty.docx")
        doc = docx.Document()
        doc.save(file_path)

        extractor = DOCXExtractor()
        with pytest.raises(EmptyDocumentError):
            extractor.extract(file_path)

    def test_extract_corrupted_docx_raises_corrupted_error(self, temp_dir):
        file_path = os.path.join(temp_dir, "corrupt.docx")
        with open(file_path, "wb") as f:
            f.write(b"CORRUPTED_NON_ZIP_DATA")

        extractor = DOCXExtractor()
        with pytest.raises(CorruptedFileError):
            extractor.extract(file_path)


class TestExtractorFactory:
    def test_get_extractor_by_extension(self):
        assert isinstance(get_extractor("report.pdf"), PDFExtractor)
        assert isinstance(get_extractor("report.docx"), DOCXExtractor)
        assert isinstance(get_extractor("readme.txt"), TXTExtractor)
        assert isinstance(get_extractor(".pdf"), PDFExtractor)

    def test_get_extractor_unsupported_type_raises_error(self):
        with pytest.raises(UnsupportedFileTypeError):
            get_extractor("archive.zip")

        with pytest.raises(UnsupportedFileTypeError):
            get_extractor("image.png")
