"""
Tests for End-to-End Vector Ingestion Pipeline.
"""

import os
import tempfile
from unittest.mock import MagicMock
import chromadb
from chromadb.config import Settings as ChromaSettings
import docx
import fitz  # PyMuPDF
import pytest

from src.document_processing.chunker import DocumentChunker
from src.embeddings.mock import MockEmbeddingService
from src.ingestion.exceptions import (
    IngestionEmbeddingError,
    IngestionExtractionError,
    IngestionStorageError,
    MissingTenantError,
)
from src.ingestion.models import IngestionRequest, IngestionResult
from src.ingestion.pipeline import IngestionPipeline
from src.vectorstore.chroma_client import verify_chroma_connection
from src.vectorstore.collection_manager import CollectionManager


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmp:
        yield tmp


@pytest.fixture
def isolated_collection_manager():
    client = chromadb.Client(ChromaSettings(is_persistent=False, anonymized_telemetry=False))
    return CollectionManager(client=client)


@pytest.fixture
def mock_embedder():
    return MockEmbeddingService(dimension=64)


@pytest.fixture
def pipeline(isolated_collection_manager, mock_embedder):
    chunker = DocumentChunker(chunk_size=100, chunk_overlap=20)
    return IngestionPipeline(
        collection_manager=isolated_collection_manager,
        embedding_service=mock_embedder,
        chunker=chunker,
        batch_size=2,
    )


class TestEndToEndIngestion:
    def test_ingest_txt_document_success(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "policy.txt")
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("Section 1: General Workplace Safety Rules.\n\nSection 2: Fire Extinguisher Locations.")

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-txt-1",
            company_id="tenant-corp-a",
            department_id="dept-safety",
            classification="internal",
            category="Policy",
            source="policy.txt",
        )

        result = pipeline.ingest_document(req)

        assert isinstance(result, IngestionResult)
        assert result.status == "success"
        assert result.document_id == "doc-txt-1"
        assert result.company_id == "tenant-corp-a"
        assert result.collection_name == "company_tenant-corp-a"
        assert result.chunks_count >= 1
        assert result.vectors_stored == result.chunks_count
        assert result.duration_ms > 0

        # Verify vectors in collection
        col = pipeline.collection_manager.get_collection("company_tenant-corp-a")
        assert col.count() == result.vectors_stored

        # Query and inspect metadata preservation
        query_res = col.get(where={"documentId": {"$eq": "doc-txt-1"}})
        assert len(query_res["ids"]) == result.vectors_stored
        for meta in query_res["metadatas"]:
            assert meta["companyId"] == "tenant-corp-a"
            assert meta["documentId"] == "doc-txt-1"
            assert meta["departmentId"] == "dept-safety"
            assert meta["classification"] == "internal"
            assert meta["category"] == "Policy"

    def test_ingest_pdf_document_with_page_provenance(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "manual.pdf")
        pdf = fitz.open()
        p1 = pdf.new_page()
        p1.insert_text((50, 50), "Page 1: Machinery Start Sequence.")
        p2 = pdf.new_page()
        p2.insert_text((50, 50), "Page 2: Emergency Shutdown Sequence.")
        pdf.save(file_path)
        pdf.close()

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-pdf-1",
            company_id="tenant-corp-b",
            department_id="dept-ops",
            source="manual.pdf",
        )

        result = pipeline.ingest_document(req)
        assert result.status == "success"
        assert result.chunks_count >= 2

        col = pipeline.collection_manager.get_collection("company_tenant-corp-b")
        res = col.get(where={"documentId": {"$eq": "doc-pdf-1"}})
        pages = [m.get("page") for m in res["metadatas"]]
        assert 1 in pages
        assert 2 in pages

    def test_ingest_docx_document_success(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "sop.docx")
        doc = docx.Document()
        doc.add_paragraph("SOP Step 1: Initialize conveyor belts.")
        doc.add_paragraph("SOP Step 2: Calibrate optical sensors.")
        doc.save(file_path)

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-docx-1",
            company_id="tenant-corp-c",
            source="sop.docx",
        )

        result = pipeline.ingest_document(req)
        assert result.status == "success"
        assert result.vectors_stored >= 1


class TestIngestionIdempotency:
    def test_reingesting_same_document_does_not_duplicate_vectors(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "idempotent.txt")
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("Line 1 text.\n\nLine 2 text.\n\nLine 3 text.")

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-idem-1",
            company_id="tenant-idem",
        )

        # Ingestion 1
        res1 = pipeline.ingest_document(req)
        col = pipeline.collection_manager.get_collection("company_tenant-idem")
        count_after_first = col.count()
        assert count_after_first == res1.vectors_stored

        # Ingestion 2 (Same document, same IDs)
        res2 = pipeline.ingest_document(req)
        count_after_second = col.count()

        # Vector count in collection must NOT double
        assert count_after_second == count_after_first
        assert res2.vectors_stored == res1.vectors_stored


class TestIngestionFailuresAndSecurity:
    def test_missing_company_id_strictly_fails(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "test.txt")
        with open(file_path, "w") as f:
            f.write("Hello world")

        with pytest.raises(MissingTenantError):
            # Bypass Pydantic validation to verify pipeline internal security check
            pipeline.ingest_document(
                IngestionRequest.model_construct(
                    file_path=file_path,
                    document_id="doc-1",
                    company_id="",
                )
            )

    def test_nonexistent_file_raises_extraction_error(self, pipeline):
        req = IngestionRequest(
            file_path="/invalid/nonexistent_file_12345.txt",
            document_id="doc-missing",
            company_id="tenant-1",
        )
        with pytest.raises(IngestionExtractionError):
            pipeline.ingest_document(req)

    def test_empty_document_raises_extraction_error(self, pipeline, temp_dir):
        file_path = os.path.join(temp_dir, "empty.txt")
        with open(file_path, "w") as f:
            f.write("   \n\n\t  ")

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-empty",
            company_id="tenant-1",
        )
        with pytest.raises(IngestionExtractionError):
            pipeline.ingest_document(req)

    def test_embedding_failure_raises_ingestion_embedding_error(self, isolated_collection_manager, temp_dir):
        file_path = os.path.join(temp_dir, "valid.txt")
        with open(file_path, "w") as f:
            f.write("Valid document text for embedding failure test.")

        failing_embedder = MagicMock()
        failing_embedder.embed_documents.side_effect = RuntimeError("Embedding model GPU OOM")
        failing_embedder.model_name = "failing-model"

        pipe = IngestionPipeline(
            collection_manager=isolated_collection_manager,
            embedding_service=failing_embedder,
        )

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-fail-embed",
            company_id="tenant-1",
        )
        with pytest.raises(IngestionEmbeddingError):
            pipe.ingest_document(req)

    def test_storage_failure_raises_ingestion_storage_error(self, mock_embedder, temp_dir):
        file_path = os.path.join(temp_dir, "valid2.txt")
        with open(file_path, "w") as f:
            f.write("Valid document text for storage failure test.")

        mock_col_mgr = MagicMock()
        mock_col_mgr.get_company_collection_name.return_value = "company_tenant-1"
        mock_col = MagicMock()
        mock_col.upsert.side_effect = RuntimeError("ChromaDB disk quota exceeded")
        mock_col_mgr.get_or_create_collection.return_value = mock_col

        pipe = IngestionPipeline(
            collection_manager=mock_col_mgr,
            embedding_service=mock_embedder,
        )

        req = IngestionRequest(
            file_path=file_path,
            document_id="doc-fail-store",
            company_id="tenant-1",
        )
        with pytest.raises(IngestionStorageError):
            pipe.ingest_document(req)


class TestLiveChromaIntegration:
    @pytest.mark.asyncio
    async def test_live_chromadb_if_available(self, temp_dir):
        """
        Integration test against live ChromaDB instance if running.
        If not reachable, gracefully skip without failing the test suite.
        """
        is_live = await verify_chroma_connection()
        if not is_live:
            pytest.skip("ChromaDB daemon is not currently reachable on localhost:8000; skipping live integration test.")

        file_path = os.path.join(temp_dir, "live_integration.txt")
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("Live ChromaDB ingestion verification test text.")

        pipe = IngestionPipeline()  # Uses live default singletons
        req = IngestionRequest(
            file_path=file_path,
            document_id="live-test-doc-001",
            company_id="live-tenant-001",
        )
        result = pipe.ingest_document(req)
        assert result.status == "success"
        assert result.vectors_stored >= 1
